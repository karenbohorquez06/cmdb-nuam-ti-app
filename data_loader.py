"""Carga de archivos CSV/Excel con mensajes de error en lenguaje natural,
selección de hoja y detección preventiva de problemas estructurales
(celdas combinadas, filas/columnas ocultas o agrupadas, encabezados vacíos...).
"""
import csv
import io
import re
import zipfile
from datetime import date, datetime, time

import pandas as pd

SUPPORTED_EXTENSIONS = {".csv", ".txt", ".xlsx", ".xlsm", ".xls"}
OPENPYXL_EXTENSIONS = {".xlsx", ".xlsm"}


class FileLoadError(Exception):
    """Error de carga con un mensaje pensado para el usuario final."""

    def __init__(self, message, suggestion=None, technical=None):
        super().__init__(message)
        self.message = message
        self.suggestion = suggestion
        self.technical = technical


def get_extension(file_name):
    name = str(file_name).lower()
    return name[name.rfind("."):] if "." in name else ""


def _as_buffer(data):
    if isinstance(data, (bytes, bytearray)):
        return io.BytesIO(data)
    if hasattr(data, "getvalue"):
        return io.BytesIO(data.getvalue())
    return data


def check_extension(file_name):
    ext = get_extension(file_name)
    if ext not in SUPPORTED_EXTENSIONS:
        shown = ext or "sin extensión"
        raise FileLoadError(
            f"El archivo «{file_name}» tiene un formato no soportado ({shown}).",
            "Use un archivo de Excel (.xlsx, .xlsm, .xls) o de texto separado por comas o punto y coma (.csv). "
            "Si el archivo viene de otra herramienta, ábralo en Excel y use «Guardar como → Libro de Excel (.xlsx)».",
        )
    return ext


# ---------------------------------------------------------------------------
# Hojas
# ---------------------------------------------------------------------------
def get_sheet_info(file_name, data):
    """Devuelve una lista de dicts {nombre, visible, filas, columnas} para archivos Excel.
    Para CSV devuelve una lista vacía."""
    ext = check_extension(file_name)
    if ext in (".csv", ".txt"):
        return []
    try:
        if ext in OPENPYXL_EXTENSIONS:
            from openpyxl import load_workbook

            wb = load_workbook(_as_buffer(data), read_only=True, data_only=True)
            info = [
                {
                    "nombre": ws.title,
                    "visible": ws.sheet_state == "visible",
                    "filas": ws.max_row or 0,
                    "columnas": ws.max_column or 0,
                }
                for ws in wb.worksheets
            ]
            wb.close()
            return info
        xls = pd.ExcelFile(_as_buffer(data))
        return [{"nombre": s, "visible": True, "filas": None, "columnas": None} for s in xls.sheet_names]
    except FileLoadError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise to_friendly_error(exc, file_name) from exc


# ---------------------------------------------------------------------------
# Lectura
# ---------------------------------------------------------------------------
def _cell_to_text(value):
    """Convierte un valor de celda a texto sin perder su formato original."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "Sí" if value else "No"
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        return repr(value)
    text = str(value)
    return text if text.strip() != "" else None


def _unique_headers(raw_headers):
    headers, seen, issues = [], {}, {"vacios": [], "duplicados": []}
    for i, h in enumerate(raw_headers, start=1):
        name = "" if h is None else str(h).strip()
        if not name:
            name = f"Columna_{i}"
            issues["vacios"].append(i)
        if name in seen:
            seen[name] += 1
            issues["duplicados"].append(name)
            name = f"{name} ({seen[name]})"
        else:
            seen[name] = 1
        headers.append(name)
    return headers, issues


def _column_letter(idx):
    from openpyxl.utils import get_column_letter

    return get_column_letter(idx)


def _read_openpyxl(data, sheet, header_row, fill_merged, skip_hidden_rows, skip_hidden_cols):
    from openpyxl import load_workbook

    wb = load_workbook(_as_buffer(data), data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    warnings = inspect_worksheet(ws, header_row)

    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    if not grid:
        raise FileLoadError(
            f"La hoja «{ws.title}» está vacía.",
            "Seleccione otra hoja que contenga datos.",
        )

    if fill_merged:
        for rng in ws.merged_cells.ranges:
            top_left = ws.cell(rng.min_row, rng.min_col).value
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    if r - 1 < len(grid) and c - 1 < len(grid[r - 1]):
                        grid[r - 1][c - 1] = top_left

    hidden_rows = {i for i, d in ws.row_dimensions.items() if d.hidden}
    hidden_cols = set()
    for d in ws.column_dimensions.values():
        if d.hidden:
            lo, hi = d.min or 0, d.max or 0
            hidden_cols.update(range(lo, hi + 1))

    if header_row > len(grid):
        raise FileLoadError(
            f"La fila de encabezado indicada ({header_row}) es mayor que el número de filas con datos ({len(grid)}).",
            "Indique una fila de encabezado más pequeña.",
        )

    ncols = max(len(r) for r in grid)
    keep_cols = [c for c in range(1, ncols + 1) if not (skip_hidden_cols and c in hidden_cols)]
    header_values = grid[header_row - 1]
    raw_headers = [header_values[c - 1] if c - 1 < len(header_values) else None for c in keep_cols]

    rows, row_numbers = [], []
    for r_idx in range(header_row + 1, len(grid) + 1):
        if skip_hidden_rows and r_idx in hidden_rows:
            continue
        src = grid[r_idx - 1]
        rows.append([_cell_to_text(src[c - 1]) if c - 1 < len(src) else None for c in keep_cols])
        row_numbers.append(r_idx)

    headers, header_issues = _unique_headers(
        [_cell_to_text(h) for h in raw_headers]
    )
    # El índice es el número de fila en Excel, para que los mensajes coincidan con lo que ve el usuario.
    df = pd.DataFrame(rows, columns=headers, index=pd.Index(row_numbers, name="Fila"), dtype="object")

    # Quitar columnas "fantasma": sin encabezado y sin ningún dato.
    ghost = [h for i, h in enumerate(headers) if (i + 1) in header_issues["vacios"] and df[h].isna().all()]
    if ghost:
        df = df.drop(columns=ghost)
        header_issues["vacios"] = [i for i in header_issues["vacios"] if headers[i - 1] not in ghost]
    wb.close()
    return df, warnings, header_issues


def _read_xls(data, sheet, header_row):
    try:
        df = pd.read_excel(_as_buffer(data), sheet_name=sheet or 0, header=header_row - 1, dtype=str)
    except ImportError as exc:
        raise FileLoadError(
            "Los archivos .xls (Excel 97-2003) necesitan un componente adicional que no está instalado.",
            "Abra el archivo en Excel y guárdelo como «Libro de Excel (.xlsx)», o instale el paquete «xlrd».",
            technical=str(exc),
        ) from exc
    headers, issues = _unique_headers([None if str(c).startswith("Unnamed:") else c for c in df.columns])
    df.columns = headers
    warnings = [{
        "nivel": "info",
        "titulo": "Formato Excel antiguo (.xls)",
        "detalle": "En archivos .xls no es posible detectar celdas combinadas ni columnas ocultas. "
                   "Para un análisis preventivo completo guarde el archivo como .xlsx.",
    }]
    return df, warnings, issues


def _decode_text(raw):
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(enc), enc
        except UnicodeDecodeError:
            continue
    raise FileLoadError(
        "No fue posible reconocer la codificación de caracteres del archivo (tildes, ñ, etc.).",
        "Guarde el archivo nuevamente como «CSV UTF-8 (delimitado por comas)» desde Excel.",
    )


def _detect_separator(text):
    sample = "\n".join(text.splitlines()[:50])
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        first = text.splitlines()[0] if text else ""
        counts = {d: first.count(d) for d in (";", ",", "\t", "|")}
        best = max(counts, key=counts.get)
        return best if counts[best] > 0 else ","


def _read_csv(data, header_row):
    raw = data.getvalue() if hasattr(data, "getvalue") else bytes(data)
    if not raw.strip():
        raise FileLoadError("El archivo está vacío.", "Verifique que seleccionó el archivo correcto.")
    text, encoding = _decode_text(raw)
    sep = _detect_separator(text)
    df = pd.read_csv(
        io.StringIO(text), sep=sep, header=header_row - 1, dtype=str,
        keep_default_na=False, na_values=[""], skipinitialspace=True,
    )
    raw_cols = [None if str(c).startswith("Unnamed:") else c for c in df.columns]
    # pandas renombra duplicados como "col.1"; los reconstruimos para avisar al usuario.
    base_cols = []
    for c in raw_cols:
        m = re.match(r"^(.*)\.(\d+)$", str(c)) if c is not None else None
        base_cols.append(m.group(1) if m and m.group(1) in raw_cols else c)
    headers, issues = _unique_headers(base_cols)
    df.columns = headers
    sep_name = {",": "coma", ";": "punto y coma", "\t": "tabulación", "|": "barra vertical"}.get(sep, sep)
    warnings = []
    if df.shape[1] == 1 and any(d in str(df.columns[0]) for d in (",", ";", "\t")):
        warnings.append({
            "nivel": "error",
            "titulo": "No se reconoció el separador de columnas",
            "detalle": "Todo el contenido quedó en una sola columna. Revise que el archivo use un único "
                       "separador (coma o punto y coma) en todas las filas.",
        })
    else:
        nombre_encoding = {"utf-8-sig": "UTF-8", "cp1252": "Windows-1252 (ANSI)", "latin-1": "ISO-8859-1"}.get(encoding, encoding)
        warnings.append({
            "nivel": "info",
            "titulo": "Detección automática",
            "detalle": f"Separador de columnas: «{sep_name}». Codificación de caracteres: {nombre_encoding}.",
        })
    return df, warnings, issues


def read_file(uploaded_file, sheet=None, header_row=1, fill_merged=False,
              skip_hidden_rows=False, skip_hidden_cols=False, return_warnings=False):
    """Lee un archivo subido (o bytes con atributo name) y devuelve un DataFrame con
    todos los valores como texto, para no perder el formato original (ceros a la
    izquierda, correos, fechas...). La conversión de tipos se hace después con
    schema_mapper.apply_schema.

    Si return_warnings=True devuelve (df, advertencias)."""
    file_name = getattr(uploaded_file, "name", "archivo")
    return load_table(file_name, uploaded_file.getvalue(), sheet, header_row, fill_merged,
                      skip_hidden_rows, skip_hidden_cols, return_warnings)


def load_table(file_name, data, sheet=None, header_row=1, fill_merged=False,
               skip_hidden_rows=False, skip_hidden_cols=False, return_warnings=True):
    ext = check_extension(file_name)
    try:
        if ext in (".csv", ".txt"):
            df, warnings, header_issues = _read_csv(data, header_row)
        elif ext in OPENPYXL_EXTENSIONS:
            df, warnings, header_issues = _read_openpyxl(
                data, sheet, header_row, fill_merged, skip_hidden_rows, skip_hidden_cols)
        else:
            df, warnings, header_issues = _read_xls(data, sheet, header_row)
    except FileLoadError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise to_friendly_error(exc, file_name) from exc

    if ext not in OPENPYXL_EXTENSIONS:
        df.index = pd.RangeIndex(header_row + 1, header_row + 1 + len(df), name="Fila")
    if df.empty and df.shape[1] == 0:
        raise FileLoadError(
            "No se encontraron columnas con datos en el archivo.",
            "Verifique que la fila de encabezado sea la correcta y que la hoja seleccionada tenga información.",
        )
    warnings = warnings + inspect_dataframe(df, header_issues)
    return (df, warnings) if return_warnings else df


# ---------------------------------------------------------------------------
# Inspección preventiva
# ---------------------------------------------------------------------------
def _fmt_list(items, limit=6):
    items = list(items)
    shown = ", ".join(str(i) for i in items[:limit])
    return shown + (f" y {len(items) - limit} más" if len(items) > limit else "")


def _ranges(nums):
    """[1,2,3,7,8] -> ['1-3', '7-8']"""
    nums = sorted(nums)
    out, start, prev = [], None, None
    for n in nums:
        if start is None:
            start = prev = n
        elif n == prev + 1:
            prev = n
        else:
            out.append(f"{start}-{prev}" if start != prev else str(start))
            start = prev = n
    if start is not None:
        out.append(f"{start}-{prev}" if start != prev else str(start))
    return out


def inspect_worksheet(ws, header_row=1):
    """Detecta celdas combinadas, filas/columnas ocultas y grupos colapsados."""
    warnings = []
    merged = list(ws.merged_cells.ranges)
    if merged:
        in_header = [str(r) for r in merged if r.min_row <= header_row <= r.max_row]
        in_data = [str(r) for r in merged if r.max_row > header_row]
        detalle = f"Se encontraron {len(merged)} rangos de celdas combinadas ({_fmt_list(str(r) for r in merged)}). "
        if in_header:
            detalle += f"Algunas afectan la fila de encabezado ({_fmt_list(in_header)}), lo que puede generar nombres de columna vacíos o repetidos. "
        if in_data:
            detalle += "En las filas de datos, solo la primera celda de cada rango conserva el valor; el resto quedará vacío "
            detalle += "a menos que active «Rellenar celdas combinadas»."
        warnings.append({"nivel": "warning", "titulo": "Celdas combinadas", "detalle": detalle})

    hidden_cols, collapsed_cols = set(), set()
    for d in ws.column_dimensions.values():
        cols = set(range(d.min or 0, (d.max or 0) + 1)) if d.min else set()
        if d.hidden:
            hidden_cols |= cols
        if d.outline_level and d.outline_level > 0:
            collapsed_cols |= cols
    if hidden_cols:
        letters = [_column_letter(c) for c in sorted(hidden_cols) if c > 0]
        warnings.append({
            "nivel": "warning",
            "titulo": "Columnas ocultas o colapsadas",
            "detalle": f"Las columnas {_fmt_list(letters)} están ocultas en Excel. Se incluirán en el análisis "
                       "a menos que active «Excluir columnas ocultas».",
        })
    elif collapsed_cols:
        letters = [_column_letter(c) for c in sorted(collapsed_cols) if c > 0]
        warnings.append({
            "nivel": "info",
            "titulo": "Columnas agrupadas",
            "detalle": f"Las columnas {_fmt_list(letters)} forman parte de un grupo (esquema) de Excel.",
        })

    hidden_rows = sorted(i for i, d in ws.row_dimensions.items() if d.hidden)
    if hidden_rows:
        warnings.append({
            "nivel": "warning",
            "titulo": "Filas ocultas, filtradas o colapsadas",
            "detalle": f"Hay {len(hidden_rows)} filas ocultas (filas {_fmt_list(_ranges(hidden_rows))}). "
                       "Puede deberse a un filtro activo o a grupos colapsados. Se incluirán en el análisis "
                       "a menos que active «Excluir filas ocultas».",
        })
    if ws.auto_filter and ws.auto_filter.ref:
        warnings.append({
            "nivel": "info",
            "titulo": "Filtro activo",
            "detalle": f"La hoja tiene un autofiltro en el rango {ws.auto_filter.ref}.",
        })
    return warnings


def inspect_dataframe(df, header_issues=None):
    """Revisiones estructurales sobre el DataFrame ya leído."""
    warnings = []
    header_issues = header_issues or {}
    if header_issues.get("vacios"):
        warnings.append({
            "nivel": "warning",
            "titulo": "Columnas sin encabezado",
            "detalle": f"Las columnas en las posiciones {_fmt_list(header_issues['vacios'])} no tienen nombre; "
                       "se nombraron automáticamente como «Columna_N». Suele indicar celdas combinadas o que "
                       "el encabezado no está en la fila indicada.",
        })
    if header_issues.get("duplicados"):
        warnings.append({
            "nivel": "warning",
            "titulo": "Encabezados repetidos",
            "detalle": f"Los nombres {_fmt_list(sorted(set(header_issues['duplicados'])))} aparecen más de una vez; "
                       "se les agregó un sufijo (2), (3)... para diferenciarlos.",
        })
    if df.empty:
        warnings.append({"nivel": "error", "titulo": "Sin registros", "detalle": "La hoja tiene encabezados pero ninguna fila de datos."})
        return warnings

    empty_cols = [c for c in df.columns if df[c].isna().all()]
    if empty_cols:
        warnings.append({
            "nivel": "warning",
            "titulo": "Columnas completamente vacías",
            "detalle": f"Las columnas {_fmt_list(empty_cols)} no contienen ningún valor.",
        })
    empty_rows = int(df.isna().all(axis=1).sum())
    if empty_rows:
        warnings.append({
            "nivel": "info",
            "titulo": "Filas vacías",
            "detalle": f"Se encontraron {empty_rows} filas totalmente vacías; se descartarán al procesar.",
        })
    numeric_headers = [c for c in df.columns if re.fullmatch(r"-?\d+([.,]\d+)?", str(c).strip())]
    if len(df.columns) and len(numeric_headers) >= max(2, len(df.columns) // 2):
        warnings.append({
            "nivel": "warning",
            "titulo": "El encabezado parece contener datos",
            "detalle": "Varios nombres de columna son números. Es posible que la primera fila no sea el encabezado; "
                       "ajuste «Fila de encabezado».",
        })
    return warnings


# ---------------------------------------------------------------------------
# Errores en lenguaje natural
# ---------------------------------------------------------------------------
def to_friendly_error(exc, file_name="el archivo"):
    """Traduce una excepción técnica a un FileLoadError con lenguaje claro."""
    if isinstance(exc, FileLoadError):
        return exc
    text = str(exc)
    tech = f"{type(exc).__name__}: {text}"

    if isinstance(exc, pd.errors.EmptyDataError):
        return FileLoadError(f"«{file_name}» está vacío o no tiene encabezados.",
                             "Verifique que el archivo tenga al menos una fila de títulos y una de datos.", tech)
    if isinstance(exc, pd.errors.ParserError):
        m = re.search(r"Expected (\d+) fields in line (\d+), saw (\d+)", text)
        if m:
            esperado, linea, visto = m.groups()
            return FileLoadError(
                f"La línea {linea} de «{file_name}» tiene {visto} columnas, pero el encabezado define {esperado}.",
                "Probablemente algún valor contiene el separador (coma o punto y coma) sin estar entre comillas. "
                "Corrija esa línea o exporte nuevamente el archivo desde Excel.", tech)
        return FileLoadError(f"La estructura de «{file_name}» no es consistente y no se pudo leer.",
                             "Revise que todas las filas tengan la misma cantidad de columnas.", tech)
    if isinstance(exc, zipfile.BadZipFile) or "zip file" in text.lower():
        return FileLoadError(
            f"«{file_name}» no es un libro de Excel válido o está dañado.",
            "Puede estar protegido con contraseña, o ser un archivo de otro tipo con extensión .xlsx. "
            "Ábralo en Excel, quite la contraseña si la tiene y guárdelo nuevamente como .xlsx.", tech)
    if isinstance(exc, KeyError) and "worksheet" in text.lower():
        return FileLoadError("La hoja seleccionada no existe en el archivo.", "Seleccione otra hoja de la lista.", tech)
    if isinstance(exc, ValueError) and "worksheet named" in text.lower():
        return FileLoadError("La hoja seleccionada no existe en el archivo.", "Seleccione otra hoja de la lista.", tech)
    if isinstance(exc, UnicodeDecodeError):
        return FileLoadError("El archivo contiene caracteres que no se pudieron interpretar.",
                             "Guárdelo como «CSV UTF-8 (delimitado por comas)».", tech)
    if isinstance(exc, MemoryError):
        return FileLoadError("El archivo es demasiado grande para procesarlo.",
                             "Divida el archivo en partes más pequeñas o elimine columnas que no necesite.", tech)
    if isinstance(exc, ImportError):
        return FileLoadError("Falta un componente para leer este tipo de archivo.",
                             "Guarde el archivo como .xlsx o .csv e intente nuevamente.", tech)
    return FileLoadError(f"No fue posible leer «{file_name}».",
                         "Verifique que el archivo no esté abierto en otro programa ni dañado, y que sea .xlsx o .csv.",
                         tech)


def friendly_message(exc, context="procesar la información"):
    """Mensaje genérico en lenguaje natural para errores fuera de la carga."""
    if isinstance(exc, FileLoadError):
        return exc.message, exc.suggestion
    if isinstance(exc, KeyError):
        return (f"No se encontró la columna {exc} necesaria para {context}.",
                "Revise que los nombres de columna coincidan con el esquema seleccionado.")
    if isinstance(exc, TypeError):
        return (f"Hay valores con un tipo de dato inesperado al {context}.",
                "Vuelva al paso «Esquema de datos» y confirme el tipo de cada columna (por ejemplo, Decimal para cantidades).")
    if isinstance(exc, MemoryError):
        return ("La información es demasiado grande para procesarla.", "Reduzca el tamaño del archivo.")
    return (f"Ocurrió un problema inesperado al {context}.",
            "Revise los archivos cargados y los tipos de datos asignados. Si persiste, comparta el detalle técnico con soporte.")
