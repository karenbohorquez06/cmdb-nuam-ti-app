"""Archivos de ejemplo que usan las pruebas. Se generan en memoria, sin depender del disco."""
import io

import pandas as pd

# Archivo 1: la clave L-1 se repite 3 veces; incluye tildes, mayúsculas y decimales con coma.
CSV_A = (
    "ID;Nombre;Responsable / Correo;Cantidad Total;Fecha Compra;Pais;Auditoria\n"
    "L-1;Camión rojo;ana@empresa.com;100,50;15/01/2024;Chile;2024-01-01 10:00\n"
    "L-1;Camión azul;ana@empresa.com;200;2024-02-01;Chile;2024-01-01 10:00\n"
    "L-1;Camión gris;ana@empresa.com;300;2024-03-01;Chile;2024-01-01 10:00\n"
    "L-2;Hola;ana@empresa.com;50;2024-04-01;Perú;2024-01-01 10:00\n"
    "L-4;Servidor A;ana@empresa.com;7;2024-06-01;Chile;2024-01-01 10:00\n"
).encode("cp1252")   # a propósito en Windows-1252, para probar la detección de codificación

# Archivo 2: la misma clave aparece 2 veces; diferencias de tilde, mayúsculas, espacios y decimal.
CSV_B = (
    "ID;Nombre;Responsable / Correo;Cantidad Total;Fecha Compra;Pais;Auditoria\n"
    "L-1;Camion rojo;ana@empresa.com;100,505;15/01/2024;Chile;2025-09-01 08:30\n"
    "L-1;Camión  azul;ana@empresa.com;200;01/02/2024;Chile;2025-09-01 08:30\n"
    "L-2;hola;ana@empresa.com;50;2024-04-01;Peru;2025-09-01 08:30\n"
    "L-5;Servidor B;ana@empresa.com;9;2024-07-01;Chile;2025-09-01 08:30\n"
).encode("utf-8")

OPCIONES_LECTURA = {"sheet": None, "header_row": 1, "fill_merged": False,
                    "skip_rows": False, "skip_cols": False}


def excel_multihoja():
    """Libro con 3 hojas, celdas combinadas, una columna oculta y una fila oculta."""
    from openpyxl import Workbook

    wb = Workbook()
    portada = wb.active
    portada.title = "Portada"
    portada["A1"] = "Reporte"

    datos = wb.create_sheet("Datos")
    datos.append(["Inventario TI", None, None, None])
    datos.merge_cells("A1:D1")                       # título combinado sobre el encabezado
    datos.append(["Serial", "Responsable", "Costo", "Grupo"])
    datos.append(["00123", "a@b.com", 10.5, "Red"])   # el cero inicial no debe perderse
    datos.append(["00124", "c@d.com", 3, None])
    datos.merge_cells("D3:D4")
    datos.column_dimensions["C"].hidden = True
    datos.row_dimensions[4].hidden = True

    wb.create_sheet("Oculta").sheet_state = "hidden"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def marco_infraestructura():
    """Una fila válida con los nombres de columna reales del esquema Infraestructura."""
    columnas = ["Nombre Servidor", "Marca", "Modelo", "Sistema Operativo", "Firmware UEFI",
                "IP Relacionada", "Serial", "Ambiente", "DataCenter", "Servicios",
                "Estado", "Descripcion"]
    fila = {c: "X" for c in columnas}
    fila.update({"IP Relacionada": "10.0.0.1", "Ambiente": "PRODUCCION", "Estado": "ACTIVO"})
    return pd.DataFrame([fila])
